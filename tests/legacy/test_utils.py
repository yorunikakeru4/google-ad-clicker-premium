"""Тесты utils.py: чтение файлов, паузы, разбор ответов 2captcha, отчёты.

Все файлы данных тесты создают сами в tmp_path, поэтому реальные queries.txt,
user_agents.txt, domains.txt и cookies.txt из репозитория не используются.
"""

import json
import platform
import random
import time
from urllib.parse import quote_plus

import openpyxl
import pytest

import utils


# --- get_random_sleep -------------------------------------------------------------


def test_get_random_sleep_stays_inside_given_range():
    for _ in range(50):
        assert 0.5 <= utils.get_random_sleep(0.5, 1.5) <= 1.5


def test_get_random_sleep_rounds_to_two_decimals(monkeypatch):
    monkeypatch.setattr(random, "uniform", lambda start, end: 1.23456789)

    assert utils.get_random_sleep(1, 2) == 1.23


def test_get_random_sleep_for_equal_bounds_returns_that_value(monkeypatch):
    monkeypatch.setattr(random, "uniform", lambda start, end: start)

    assert utils.get_random_sleep(2, 2) == 2.0


def test_get_random_sleep_uses_both_bounds(monkeypatch):
    seen = {}

    def fake_uniform(start, end):
        seen["range"] = (start, end)
        return start

    monkeypatch.setattr(random, "uniform", fake_uniform)

    utils.get_random_sleep(3, 9)

    assert seen["range"] == (3, 9)


# --- _check_error -----------------------------------------------------------------


@pytest.mark.parametrize(
    "response_text",
    [
        "ERROR_WRONG_USER_KEY",
        "ERROR_KEY_DOES_NOT_EXIST",
        "ERROR_ZERO_BALANCE",
        "IP_BANNED",
        "ERROR_GOOGLEKEY",
    ],
)
def test_check_error_in_php_exits_on_fatal_errors(response_text):
    exit_, cont, brk = utils._check_error(response_text)

    assert exit_ is True
    assert cont is False
    assert brk is False


def test_check_error_in_php_continues_when_no_slots_available(monkeypatch):
    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.0)

    exit_, cont, brk = utils._check_error("ERROR_NO_SLOT_AVAILABLE")

    assert (exit_, cont, brk) == (False, True, False)


def test_check_error_in_php_breaks_on_task_id(monkeypatch):
    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.0)

    exit_, cont, brk = utils._check_error("OK|2122988149")

    assert (exit_, cont, brk) == (False, False, True)


def test_check_error_in_php_fatal_error_wins_over_continue_error(monkeypatch):
    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.0)

    exit_, cont, brk = utils._check_error("ERROR_ZERO_BALANCE ERROR_NO_SLOT_AVAILABLE")

    assert (exit_, cont, brk) == (True, False, False)


def test_check_error_res_php_exits_on_wrong_user_key(monkeypatch):
    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.0)

    assert utils._check_error("ERROR_WRONG_USER_KEY", "res_php") == (True, False, False)


def test_check_error_res_php_exits_on_unsolvable_captcha(monkeypatch):
    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.0)

    assert utils._check_error("ERROR_CAPTCHA_UNSOLVABLE", "res_php") == (True, False, False)


def test_check_error_res_php_continues_while_captcha_not_ready(monkeypatch):
    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.0)

    assert utils._check_error("CAPCHA_NOT_READY", "res_php") == (False, True, False)


def test_check_error_res_php_breaks_on_solved_captcha(monkeypatch):
    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.0)

    assert utils._check_error("OK|abc123", "res_php") == (False, False, True)


def test_check_error_unknown_request_type_returns_no_flags(caplog):
    assert utils._check_error("OK|1", "unknown_type") == (False, False, False)

    # Текст запроса ушёл в поля, а не в message: проверяем структурную
    # запись в зеркале legacy-логгера, а не склеенную строку.
    record = next(
        entry for entry in caplog.records if entry.getMessage().startswith("Wrong request type")
    )
    assert record.category == "captcha"
    assert record.fields == {"request_type": "unknown_type"}


def test_check_error_in_php_treats_res_php_captcha_not_ready_as_success(monkeypatch):
    # Ошибка res.php не обрабатывается в режиме in.php и считается успехом:
    # это текущее поведение, которое важно не потерять при рефакторинге.
    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.0)

    assert utils._check_error("CAPCHA_NOT_READY") == (False, False, True)


def test_check_error_wait_is_scaled_by_wait_factor(monkeypatch):
    # error_wait = 5 * wait_factor; две точки фиксируют коэффициент 5.
    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.08)
    started = time.monotonic()
    utils._check_error("ERROR_NO_SLOT_AVAILABLE")
    assert time.monotonic() - started >= 0.4

    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.2)
    started = time.monotonic()
    utils._check_error("CAPCHA_NOT_READY", "res_php")
    assert time.monotonic() - started >= 1.0


def test_check_error_does_not_wait_when_wait_factor_is_zero(monkeypatch):
    monkeypatch.setattr(utils.config.behavior, "wait_factor", 0.0)

    started = time.monotonic()
    utils._check_error("ERROR_NO_SLOT_AVAILABLE")
    elapsed = time.monotonic() - started

    assert elapsed < 0.1


# --- Чтение списков из файлов ----------------------------------------------------


def test_get_queries_strips_whitespace_and_quotes(set_paths, tmp_path):
    query_file = tmp_path / "queries.txt"
    query_file.write_text("  'wireless keyboard'  \n\"bluetooth headphones\"\n", "utf-8")
    set_paths(query_file=query_file)

    assert utils.get_queries() == ["wireless keyboard", "bluetooth headphones"]


def test_get_queries_skips_blank_lines(set_paths, tmp_path):
    # Пустая строка в файле запросов стала бы пустым поисковым запросом:
    # кликер бы искал по пустой строке и впустую тратил трафик.
    query_file = tmp_path / "queries.txt"
    query_file.write_text("usb hub\n\n\nwebcam\n", "utf-8")
    set_paths(query_file=query_file)

    assert utils.get_queries() == ["usb hub", "webcam"]


def test_get_queries_skips_lines_of_whitespace_and_quotes(set_paths, tmp_path):
    query_file = tmp_path / "queries.txt"
    query_file.write_text("usb hub\n   \n''\n\"\"\nwebcam\n", "utf-8")
    set_paths(query_file=query_file)

    assert utils.get_queries() == ["usb hub", "webcam"]


def test_get_queries_returns_empty_list_when_file_has_only_blank_lines(set_paths, tmp_path):
    query_file = tmp_path / "queries.txt"
    query_file.write_text("\n\n \n", "utf-8")
    set_paths(query_file=query_file)

    assert utils.get_queries() == []


def test_get_domains_skips_blank_lines(set_paths, tmp_path):
    domains_file = tmp_path / "domains.txt"
    domains_file.write_text("www.booking.com\n\n  \nwww.sixt.com\n", "utf-8")
    set_paths(filtered_domains=domains_file)

    # Нормализация до хоста без www: список — чёрный, а матчинг сравнивает
    # хосты, поэтому «www.booking.com» и «booking.com» — один и тот же домен.
    assert utils.get_domains() == ["booking.com", "sixt.com"]


def test_get_domains_normalizes_urls_and_dedupes(set_paths, tmp_path):
    domains_file = tmp_path / "domains.txt"
    domains_file.write_text(
        "https://www.edelind.de\nedelind.de/\nwww.edelind.de\nEDelind.DE\n", "utf-8"
    )
    set_paths(filtered_domains=domains_file)

    assert utils.get_domains() == ["edelind.de"]


def test_get_domains_appends_own_domain(set_paths, tmp_path, monkeypatch, behavior):
    domains_file = tmp_path / "domains.txt"
    domains_file.write_text("booking.com\n", "utf-8")
    set_paths(filtered_domains=domains_file)
    monkeypatch.setattr(behavior, "own_domain", "https://www.edelind.de/goldkette")

    assert utils.get_domains() == ["booking.com", "edelind.de"]


def test_get_domains_own_domain_works_with_empty_file(set_paths, tmp_path, monkeypatch, behavior):
    domains_file = tmp_path / "domains.txt"
    domains_file.write_text("\n", "utf-8")
    set_paths(filtered_domains=domains_file)
    monkeypatch.setattr(behavior, "own_domain", "edelind.de")

    assert utils.get_domains() == ["edelind.de"]


def test_get_domains_own_domain_duplicates_file_entry(set_paths, tmp_path, monkeypatch, behavior):
    domains_file = tmp_path / "domains.txt"
    domains_file.write_text("www.edelind.de\n", "utf-8")
    set_paths(filtered_domains=domains_file)
    monkeypatch.setattr(behavior, "own_domain", "edelind.de")

    assert utils.get_domains() == ["edelind.de"]


def test_get_domains_returns_empty_list_when_file_has_only_blank_lines(set_paths, tmp_path):
    domains_file = tmp_path / "domains.txt"
    domains_file.write_text("\n''\n", "utf-8")
    set_paths(filtered_domains=domains_file)

    assert utils.get_domains() == []


def test_get_user_agents_skips_blank_lines(tmp_path):
    ua_file = tmp_path / "user_agents.txt"
    ua_file.write_text("Mozilla/5.0 Windows\n\n  \nMozilla/5.0 Linux\n", "utf-8")

    assert utils._get_user_agents(ua_file) == ["Mozilla/5.0 Windows", "Mozilla/5.0 Linux"]


def test_get_user_agents_returns_empty_list_when_file_has_only_blank_lines(tmp_path):
    ua_file = tmp_path / "user_agents.txt"
    ua_file.write_text("\n\"\"\n", "utf-8")

    assert utils._get_user_agents(ua_file) == []


def test_get_queries_exits_when_file_is_missing(set_paths, tmp_path):
    set_paths(query_file=tmp_path / "nope.txt")

    with pytest.raises(SystemExit) as excinfo:
        utils.get_queries()

    assert "Couldn't find queries file" in str(excinfo.value)


def test_get_domains_strips_whitespace_and_quotes(set_paths, tmp_path):
    domains_file = tmp_path / "domains.txt"
    domains_file.write_text(" 'www.booking.com' \n\"www.sixt.com\"\n", "utf-8")
    set_paths(filtered_domains=domains_file)

    assert utils.get_domains() == ["booking.com", "sixt.com"]


def test_get_domains_exits_when_file_is_missing(set_paths, tmp_path):
    set_paths(filtered_domains=tmp_path / "nope.txt")

    with pytest.raises(SystemExit) as excinfo:
        utils.get_domains()

    assert "Couldn't find domains file" in str(excinfo.value)


# --- domain_matches: матчинг хоста чёрного списка ---------------------------------


@pytest.mark.parametrize(
    ("url", "domain", "expected"),
    [
        # апекс и поддомен — одно совпадение
        ("https://edelind.de/goldkette", "edelind.de", True),
        ("https://www.edelind.de/goldkette", "edelind.de", True),
        ("https://shop.edelind.de/x?a=1", "edelind.de", True),
        # www в записи списка тоже снимается
        ("https://edelind.de/", "www.edelind.de", True),
        # подстрока хоста не совпадает: not-edelind.de ≠ edelind.de
        ("https://notedelind.de/", "edelind.de", False),
        ("https://edelind.de.evil.example/", "edelind.de", False),
        # домен в query/fragment — это не хост
        ("https://example.com/?next=edelind.de", "edelind.de", False),
        ("https://example.com/#edelind.de", "edelind.de", False),
        # порт и raw-строка без схемы (такие строки лежат в файле)
        ("https://edelind.de:8443/x", "edelind.de", True),
        ("edelind.de/goldkette", "edelind.de", True),
        ("www.edelind.de", "edelind.de", True),
        # пустые значения ничего не блокируют
        ("https://edelind.de/", "", False),
        ("", "edelind.de", False),
        (None, "edelind.de", False),
    ],
)
def test_domain_matches_compares_hosts_not_substrings(url, domain, expected):
    assert utils.domain_matches(url, domain) is expected


def test_get_user_agents_strips_whitespace_and_quotes(tmp_path):
    ua_file = tmp_path / "user_agents.txt"
    ua_file.write_text(" 'Mozilla/5.0 Windows' \n\"Mozilla/5.0 Linux\"\n", "utf-8")

    assert utils._get_user_agents(ua_file) == ["Mozilla/5.0 Windows", "Mozilla/5.0 Linux"]


def test_get_user_agents_exits_when_file_is_missing(tmp_path):
    with pytest.raises(SystemExit) as excinfo:
        utils._get_user_agents(tmp_path / "nope.txt")

    assert "Couldn't find user agents file" in str(excinfo.value)


# --- get_locale_language ----------------------------------------------------------


def test_get_locale_language_returns_primary_locale_for_known_country():
    assert utils.get_locale_language("US") == "en-US"


def test_get_locale_language_returns_first_locale_for_multilingual_country():
    assert utils.get_locale_language("AF") == "ps-AF"


def test_get_locale_language_returns_none_for_unknown_country():
    # None, а не прежнее "en": одиночный Accept-Language: en — редкий
    # паттерн, который на проводе читается как аномалия (отчёт §4.1).
    # Без локали браузер шлёт свой честный дефолт en-US,en;q=0.9.
    assert utils.get_locale_language("ZZ") is None


def test_get_locale_language_returns_none_without_country():
    # Геолокация упала (country_code=None) — локаль не ставится вовсе.
    assert utils.get_locale_language(None) is None


def test_get_locale_language_falls_back_to_module_dir_when_cwd_has_no_file(
    isolated_cwd,
):
    # Упакованный запуск: cwd — каталог данных без этого файла, локали
    # обязаны браться из каталога модуля (в frozen это _MEIPASS, где файл
    # лежит в бандле sidecar).
    (isolated_cwd / "country_to_locale.json").unlink()

    assert utils.get_locale_language("US") == "en-US", (
        "значения должны прийти из настоящего country_to_locale.json рядом с utils.py"
    )


def test_get_locale_language_raises_when_mapping_file_is_missing(
    isolated_cwd, monkeypatch
):
    # «Пропал» в обоих местах: нет в cwd и нет рядом с модулем — иначе
    # фолбэк спрятал бы отсутствие файла, которое здесь и проверяем.
    import engine.diagnostics as diagnostics

    (isolated_cwd / "country_to_locale.json").unlink()
    monkeypatch.setattr(
        diagnostics, "COUNTRY_LOCALES_FILE", isolated_cwd / "absent.json"
    )

    with pytest.raises(FileNotFoundError):
        utils.get_locale_language("US")


# --- Direction --------------------------------------------------------------------


def test_direction_members_keep_their_values():
    assert utils.Direction.UP.value == "UP"
    assert utils.Direction.DOWN.value == "DOWN"
    assert utils.Direction.LEFT.value == "LEFT"
    assert utils.Direction.RIGHT.value == "RIGHT"
    assert utils.Direction.BOTH.value == "BOTH"


# --- resolve_redirect -------------------------------------------------------------


@pytest.mark.parametrize(
    "broken_url",
    ["htp:/broken-url", "not-a-url", "://missing-scheme"],
)
def test_resolve_redirect_returns_original_url_when_request_fails(broken_url, caplog):
    # Сетевых обращений здесь нет: requests падает на разборе URL до выхода в сеть.
    assert utils.resolve_redirect(broken_url) == broken_url
    assert "Error resolving URL redirection" in caplog.text


# --- add_cookies ------------------------------------------------------------------


def test_add_cookies_normalizes_samesite_values(isolated_cwd, recording_cookie_driver):
    (isolated_cwd / "cookies.txt").write_text(
        json.dumps(
            [
                {"name": "a", "sameSite": "strict", "secure": True},
                {"name": "b", "sameSite": "lax", "secure": False},
                {"name": "c", "sameSite": "no_restriction", "secure": True},
                {"name": "d", "sameSite": "anything", "secure": False},
            ]
        ),
        encoding="utf-8",
    )

    utils.add_cookies(recording_cookie_driver)

    same_site = {cookie["name"]: cookie["sameSite"] for cookie in recording_cookie_driver.cookies}
    assert same_site == {"a": "Strict", "b": "Lax", "c": "None", "d": "Lax"}


def test_add_cookies_keeps_all_other_cookie_fields(isolated_cwd, recording_cookie_driver):
    (isolated_cwd / "cookies.txt").write_text(
        json.dumps(
            [{"name": "session", "value": "42", "domain": ".example.com", "sameSite": "lax"}]
        ),
        encoding="utf-8",
    )

    utils.add_cookies(recording_cookie_driver)

    assert recording_cookie_driver.cookies[0]["name"] == "session"
    assert recording_cookie_driver.cookies[0]["value"] == "42"
    assert recording_cookie_driver.cookies[0]["domain"] == ".example.com"


@pytest.mark.xfail(
    reason=(
        "Баг legacy-кода: add_cookies обращается к cookie['sameSite'] и "
        "cookie['secure'] напрямую, поэтому cookie без этих полей (например "
        "выгруженный не из Chrome) роняет кликер с голым KeyError"
    ),
    strict=True,
)
def test_add_cookies_should_tolerate_cookie_without_samesite(
    isolated_cwd, recording_cookie_driver
):
    (isolated_cwd / "cookies.txt").write_text(
        json.dumps([{"name": "session", "value": "42", "domain": ".example.com"}]),
        encoding="utf-8",
    )

    utils.add_cookies(recording_cookie_driver)

    assert recording_cookie_driver.cookies[0]["name"] == "session"


def test_add_cookies_exits_when_cookies_file_is_missing(isolated_cwd, recording_cookie_driver):
    (isolated_cwd / "cookies.txt").unlink()

    with pytest.raises(SystemExit) as excinfo:
        utils.add_cookies(recording_cookie_driver)

    assert "Missing cookies.txt file!" in str(excinfo.value)


def test_add_cookies_exits_on_malformed_json(isolated_cwd, recording_cookie_driver, caplog):
    (isolated_cwd / "cookies.txt").write_text("{oops", encoding="utf-8")

    with pytest.raises(SystemExit):
        utils.add_cookies(recording_cookie_driver)

    assert "Failed to read cookies file. Check format and try again." in caplog.text


# --- generate_click_report --------------------------------------------------------


def test_generate_click_report_writes_xlsx_with_headers(isolated_cwd):
    results = [("http://example.com", 3, "Ad", "10:11:12", "usb hub")]

    utils.generate_click_report(results, "28-09-2026")

    sheet = openpyxl.load_workbook(isolated_cwd / "click_report_28-09-2026.xlsx").active
    headers = [sheet["A1"].value, sheet["B1"].value, sheet["C1"].value, sheet["D1"].value,
               sheet["E1"].value]
    assert headers == ["URL", "Query", "Clicks", "Time", "Category"]


def test_generate_click_report_orders_columns_url_query_clicks_time_category(isolated_cwd):
    results = [("http://example.com", 3, "Ad", "10:11:12", "usb hub")]

    utils.generate_click_report(results, "28-09-2026")

    sheet = openpyxl.load_workbook(isolated_cwd / "click_report_28-09-2026.xlsx").active
    assert [sheet["A2"].value, sheet["B2"].value, sheet["C2"].value, sheet["D2"].value,
            sheet["E2"].value] == [
        "http://example.com",
        "usb hub",
        3,
        "28-09-2026 10:11:12",
        "Ad",
    ]


def test_generate_click_report_writes_one_row_per_result(isolated_cwd):
    results = [
        ("http://a.example", 1, "Ad", "10:00:00", "usb hub"),
        ("http://b.example", 2, "Non-ad", "11:00:00", "webcam"),
        ("http://c.example", 4, "Shopping", "12:00:00", "monitor"),
    ]

    utils.generate_click_report(results, "28-09-2026")

    sheet = openpyxl.load_workbook(isolated_cwd / "click_report_28-09-2026.xlsx").active
    assert sheet.max_row == 4
    assert [row[4] for row in sheet.iter_rows(min_row=2, values_only=True)] == [
        "Ad",
        "Non-ad",
        "Shopping",
    ]


def test_generate_click_report_creates_only_headers_without_results(isolated_cwd):
    utils.generate_click_report([], "28-09-2026")

    sheet = openpyxl.load_workbook(isolated_cwd / "click_report_28-09-2026.xlsx").active
    assert sheet.max_row == 1
    assert sheet["A1"].value == "URL"


def test_generate_click_report_sets_column_widths(isolated_cwd):
    utils.generate_click_report([], "28-09-2026")

    sheet = openpyxl.load_workbook(isolated_cwd / "click_report_28-09-2026.xlsx").active
    widths = {letter: sheet.column_dimensions[letter].width for letter in "ABCDE"}
    assert widths == {"A": 80, "B": 25, "C": 15, "D": 20, "E": 15}


# --- get_random_user_agent_string -------------------------------------------------


def test_get_random_user_agent_string_filters_by_linux(set_paths, tmp_path, monkeypatch):
    ua_file = tmp_path / "user_agents.txt"
    ua_file.write_text("Mozilla/5.0 (X11; Linux x86_64) Chrome/118\n"
                       "Mozilla/5.0 (Windows NT 10.0) Chrome/120\n", encoding="utf-8")
    set_paths(user_agents=ua_file)
    monkeypatch.setattr(platform, "system", lambda: "Linux")

    for _ in range(20):
        assert utils.get_random_user_agent_string() == "Mozilla/5.0 (X11; Linux x86_64) Chrome/118"


def test_get_random_user_agent_string_filters_by_windows(set_paths, tmp_path, monkeypatch):
    ua_file = tmp_path / "user_agents.txt"
    ua_file.write_text("Mozilla/5.0 (X11; Linux x86_64) Chrome/118\n"
                       "Mozilla/5.0 (Windows NT 10.0) Chrome/120\n", encoding="utf-8")
    set_paths(user_agents=ua_file)
    monkeypatch.setattr(platform, "system", lambda: "Windows")

    assert utils.get_random_user_agent_string() == "Mozilla/5.0 (Windows NT 10.0) Chrome/120"


def test_get_random_user_agent_string_drops_ios_agents_on_macos(set_paths, tmp_path, monkeypatch):
    ua_file = tmp_path / "user_agents.txt"
    ua_file.write_text("Mozilla/5.0 (iPhone; CPU iPhone OS 17_0) Safari\n"
                       "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) Chrome/131.0.0.0\n"
                       "Mozilla/5.0 (Windows NT 10.0) Chrome/120\n", encoding="utf-8")
    set_paths(user_agents=ua_file)
    monkeypatch.setattr(platform, "system", lambda: "Darwin")

    # Мобильный UA в десктопном Chrome — готовый сигнал детекта: платформа
    # MacIntel против строки iOS (см. suspicion_flags в диагностике и
    # Sec-CH-UA-Platform на стороне сервера). На macOS остаётся только Macintosh.
    assert utils.get_random_user_agent_string().startswith("Mozilla/5.0 (Macintosh")


def test_get_random_user_agent_string_drops_android_agents_on_linux(set_paths, tmp_path, monkeypatch):
    ua_file = tmp_path / "user_agents.txt"
    ua_file.write_text("Mozilla/5.0 (Linux; Android 13) Chrome/117\n"
                       "Mozilla/5.0 (X11; Linux x86_64) Chrome/131.0.0.0\n", encoding="utf-8")
    set_paths(user_agents=ua_file)
    monkeypatch.setattr(platform, "system", lambda: "Linux")

    # та же логика: Android-UA на десктопном Chrome не должен выбираться
    assert utils.get_random_user_agent_string() == "Mozilla/5.0 (X11; Linux x86_64) Chrome/131.0.0.0"


def test_get_random_user_agent_string_falls_back_to_all_agents_on_unknown_os(
    set_paths, tmp_path, monkeypatch
):
    ua_file = tmp_path / "user_agents.txt"
    ua_file.write_text("Mozilla/5.0 (Windows NT 10.0) Chrome/120\n", encoding="utf-8")
    set_paths(user_agents=ua_file)
    monkeypatch.setattr(platform, "system", lambda: "Plan9")

    assert utils.get_random_user_agent_string() == "Mozilla/5.0 (Windows NT 10.0) Chrome/120"


def test_get_random_user_agent_string_exits_when_no_agent_matches_os(
    set_paths, tmp_path, monkeypatch
):
    ua_file = tmp_path / "user_agents.txt"
    ua_file.write_text("Mozilla/5.0 (Windows NT 10.0) Chrome/120\n", encoding="utf-8")
    set_paths(user_agents=ua_file)
    monkeypatch.setattr(platform, "system", lambda: "Linux")

    with pytest.raises(IndexError):
        utils.get_random_user_agent_string()


@pytest.mark.xfail(
    reason=(
        "Баг legacy-кода: если в user_agents.txt нет строки под текущую ОС, "
        "random.choice() падает с голым IndexError без подсказки, что нужно "
        "добавить user-agent'ы для своей ОС"
    ),
    strict=True,
)
def test_get_random_user_agent_string_should_exit_with_readable_error(
    set_paths, tmp_path, monkeypatch, caplog
):
    ua_file = tmp_path / "user_agents.txt"
    ua_file.write_text("Mozilla/5.0 (Windows NT 10.0) Chrome/120\n", encoding="utf-8")
    set_paths(user_agents=ua_file)
    monkeypatch.setattr(platform, "system", lambda: "Linux")

    with pytest.raises(SystemExit):
        utils.get_random_user_agent_string()

    assert "user agents" in caplog.text.lower()


# --- is_mobile_user_agent ----------------------------------------------------------


def test_is_mobile_user_agent_detects_mobile_strings():
    assert utils.is_mobile_user_agent(
        "Mozilla/5.0 (Linux; Android 13; SM-S901B) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/112.0.0.0 Mobile Safari/537.36"
    )
    assert utils.is_mobile_user_agent(
        "Mozilla/5.0 (iPhone; CPU iPhone OS 17_7 like Mac OS X) AppleWebKit/605.1.15 "
        "(KHTML, like Gecko) CriOS/129.0.6668.69 Mobile/15E148 Safari/604.1"
    )
    assert utils.is_mobile_user_agent(
        "Mozilla/5.0 (iPad; CPU OS 17_5 like Mac OS X) AppleWebKit/605.1.15 "
        "(KHTML, like Gecko) CriOS/125.0.6422.80 Mobile/15E148 Safari/604.1"
    )


def test_is_mobile_user_agent_is_false_for_desktop_strings():
    assert not utils.is_mobile_user_agent(
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/142.0.0.0 Safari/537.36"
    )
    assert not utils.is_mobile_user_agent(
        "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36"
    )
    assert not utils.is_mobile_user_agent(
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/136.0.0.0 Safari/537.36"
    )


# --- probe_proxy_captcha ----------------------------------------------------------


@pytest.fixture(autouse=True)
def _probe_cache_isolated():
    """Кэш probe живёт в модуле между тестами — чистим до и после каждого.

    Без этого тест, проверяющий «чистую» страницу, увидел бы капчу,
    закэшированную предыдущим тестом тем же прокси.
    """

    utils._probe_cache.clear()
    yield
    utils._probe_cache.clear()


class _ProbeClock:
    """Управляемое время для TTL-кэша probe: подменяет time.monotonic в utils."""

    def __init__(self, start: float = 1000.0):
        self.now = start

    def advance(self, seconds: float) -> None:
        self.now += seconds


@pytest.fixture
def probe_clock(monkeypatch):
    """Часы для TTL-кэша: тесты сдвигают время сами, без реальных ожиданий."""

    clock = _ProbeClock()
    monkeypatch.setattr(utils, "monotonic", lambda: clock.now)
    return clock


@pytest.fixture
def probe_http(monkeypatch):
    """Сетевые вызовы probe: каждый запрос записывается, ответ — чистый."""

    calls = []

    def _get(url, proxies=None, timeout=None, headers=None):
        calls.append({"url": url, "proxies": proxies, "timeout": timeout, "headers": headers})
        return _FakeResponse(url, "<html>results</html>")

    monkeypatch.setattr(utils.requests, "get", _get)
    return calls


class _FakeResponse:
    """Ответ requests для probe: url + тело, без сети."""

    def __init__(self, url, text="", status_code=200):
        self.url = url
        self.text = text
        self.status_code = status_code


def test_probe_flags_captcha_by_sorry_url(monkeypatch):
    monkeypatch.setattr(
        utils.requests,
        "get",
        lambda *a, **k: _FakeResponse(
            "https://www.google.com/sorry/index?continue=x", "whatever"
        ),
    )

    assert utils.probe_proxy_captcha("user:pass@host:8080") is True


def test_probe_flags_captcha_by_body_text(monkeypatch):
    monkeypatch.setattr(
        utils.requests,
        "get",
        lambda *a, **k: _FakeResponse(
            "https://www.google.com/search?q=probe",
            "Our systems have detected unusual traffic from your computer network",
        ),
    )

    assert utils.probe_proxy_captcha("host:8080") is True


def test_probe_passes_clean_search_page(monkeypatch):
    monkeypatch.setattr(
        utils.requests,
        "get",
        lambda *a, **k: _FakeResponse(
            "https://www.google.com/search?q=probe", "<html>results</html>"
        ),
    )

    assert utils.probe_proxy_captcha("host:8080") is False


def test_probe_returns_none_on_network_error(monkeypatch):
    # Сетевая ошибка — None, и вызывающий обязан НЕ поднимать браузер:
    # проверка ничего не подтвердила, а Chrome на нерабочем прокси сжёг бы
    # время и трафик вхолостую.
    def _raise(*args, **kwargs):
        raise utils.requests.RequestException("boom")

    monkeypatch.setattr(utils.requests, "get", _raise)

    assert utils.probe_proxy_captcha("host:8080") is None


def test_probe_treats_error_status_as_not_clean(monkeypatch):
    # 4xx/5xx — ответ получен, но IP он не проверяет: прокси отдал 502 на
    # CONNECT или 403 на сам запрос, Google — 429/503. В body такого ответа
    # нет маркера капчи, и раньше он читался бы как «чисто» — раунд стартовал
    # с заведомо нерабочим прокси.
    for status in (403, 407, 429, 500, 502, 503):
        monkeypatch.setattr(
            utils.requests,
            "get",
            lambda *a, **k: _FakeResponse(
                "https://www.google.com/search?q=probe", "<html>error</html>", status
            ),
        )

        assert utils.probe_proxy_captcha("host:8080") is None, f"статус {status}"


def test_probe_error_status_does_not_hide_a_captcha(monkeypatch):
    # Капча, приехавшая вместе с 4xx: /sorry/ важнее кода ответа, раунд
    # всё равно пропускается — но уже как капча, с ротацией по этой причине.
    monkeypatch.setattr(
        utils.requests,
        "get",
        lambda *a, **k: _FakeResponse(
            "https://www.google.com/sorry/index?continue=x", "whatever", 429
        ),
    )

    assert utils.probe_proxy_captcha("host:8080") is True


def test_probe_passes_proxy_credentials_to_requests(monkeypatch):
    captured = {}

    def _get(url, proxies=None, timeout=None, headers=None):
        captured["proxies"] = proxies
        return _FakeResponse(url, "")

    monkeypatch.setattr(utils.requests, "get", _get)

    utils.probe_proxy_captcha("user:pass@proxy.host:8080")

    assert captured["proxies"] == {
        "http": "http://user:pass@proxy.host:8080",
        "https": "http://user:pass@proxy.host:8080",
    }


# --- probe_proxy_captcha: кэш и рандомизация ---------------------------------------


def test_second_probe_within_ttl_is_served_from_cache(probe_http, probe_clock):
    """Повторный probe того же прокси внутри TTL не ходит в сеть.

    Раунд идёт каждые ~2 минуты, поэтому подтверждённый результат живёт
    дольше одного раунда: на сессию уходит один health-check, а не один
    на каждый — иначе probe удваивает нагрузку на exit IP.
    """
    assert utils.probe_proxy_captcha("host:8080") is False

    probe_clock.advance(utils._PROBE_CACHE_TTL - 1)

    assert utils.probe_proxy_captcha("host:8080") is False
    assert len(probe_http) == 1, "внутри TTL второй probe обязан взять кэш"


def test_probe_goes_to_network_again_after_ttl_expires(probe_http, probe_clock):
    """После истечения TTL probe снова спрашивает сеть, а не кэш вечно."""
    assert utils.probe_proxy_captcha("host:8080") is False

    probe_clock.advance(utils._PROBE_CACHE_TTL)

    assert utils.probe_proxy_captcha("host:8080") is False
    assert len(probe_http) == 2, "просроченный кэш не должен отдавать результат"


def test_probe_cache_is_isolated_per_proxy(probe_http, probe_clock):
    """Разные прокси кэшируются раздельно: кэш одного не закрывает другой."""
    assert utils.probe_proxy_captcha("host-a:8080") is False
    assert utils.probe_proxy_captcha("host-b:8080") is False
    assert len(probe_http) == 2, "каждая строка прокси проверяется своей сетью"

    assert utils.probe_proxy_captcha("host-a:8080") is False

    assert len(probe_http) == 2, "кэш host-a не должен пострадать от probe host-b"


def test_probe_caches_captcha_verdict(monkeypatch, probe_clock):
    """Капча тоже кэшируется: повторный probe не спрашивает сеть о мёртвом IP."""
    calls = []

    def _get(url, **kwargs):
        calls.append(url)
        return _FakeResponse("https://www.google.com/sorry/index?continue=x", "")

    monkeypatch.setattr(utils.requests, "get", _get)

    assert utils.probe_proxy_captcha("host:8080") is True

    probe_clock.advance(utils._PROBE_CACHE_TTL - 1)

    assert utils.probe_proxy_captcha("host:8080") is True
    assert len(calls) == 1


def test_unknown_probe_result_is_not_cached(monkeypatch):
    """None — «проверить не удалось»: транзиентный обрыв не кэшируется.

    Иначе один сбой сети заблокировал бы раунды на весь TTL, хотя через
    секунды прокси мог ожить.
    """
    attempts = []

    def _get(url, **kwargs):
        attempts.append(url)
        if len(attempts) == 1:
            raise utils.requests.RequestException("boom")
        return _FakeResponse(url, "<html>results</html>")

    monkeypatch.setattr(utils.requests, "get", _get)

    assert utils.probe_proxy_captcha("host:8080") is None
    assert utils.probe_proxy_captcha("host:8080") is False
    assert len(attempts) == 2, "неудачный probe обязан повторить запрос"


def test_probe_query_is_randomized():
    """Probe берёт разные безобидные query, а не одну фиксированную строку.

    Постоянный «proxy+health+check» с одного IP перед каждым раундом —
    сигнатура бота; вариативность проверяется на большом числе выборок,
    чтобы совпадение первого запроса не дало ложный провал.
    """
    urls = [utils._probe_url() for _ in range(30)]

    assert len(set(urls)) > 1, "query обязан меняться от запроса к запросу"

    queries = {url.split("q=", 1)[1] for url in urls}
    expected = {quote_plus(item) for item in utils._PROBE_QUERIES}
    assert queries <= expected, "probe не должен уходить за список безобидных query"


def test_probe_request_uses_benign_query_and_chrome_ua(probe_http):
    """Сетевой запрос probe: бытовой query и честный Chrome-UA из диапазона."""
    utils.probe_proxy_captcha("host:8080")

    request = probe_http[0]
    query = request["url"].split("q=", 1)[1]
    assert query in {quote_plus(item) for item in utils._PROBE_QUERIES}

    user_agent = request["headers"]["User-Agent"]
    assert user_agent.startswith("Mozilla/5.0 (Windows NT 10.0; Win64; x64)")
    version = int(user_agent.split("Chrome/")[1].split(".")[0])
    assert version in utils._PROBE_UA_VERSIONS, "версия UA должна варьироваться в диапазоне"


# --- require_german_exit ----------------------------------------------------------


class TestRequireGermanExit:
    """Гейт «только Германия»: срабатывает до старта Chrome, без кредов в тексте."""

    @pytest.mark.parametrize("code", ["DE", "de", " dE "])
    def test_germany_passes(self, code):
        assert utils.require_german_exit(code, "user:s3cr3t@proxy.host:8080") == "DE"

    @pytest.mark.parametrize("code", ["TR", "FR", "US", "ke"])
    def test_any_other_country_blocks(self, code):
        with pytest.raises(utils.ProxyCountryError) as excinfo:
            utils.require_german_exit(code, "user:s3cr3t@proxy.host:8080")

        message = str(excinfo.value)
        assert code.upper() in message
        assert "proxy.host:8080" in message, "адрес нужен для причины в логе"
        # Креды не должны дойти ни до лога, ни до workers.last_error.
        assert "s3cr3t" not in message
        assert "user:" not in message

    @pytest.mark.parametrize("code", [None, "", "   "])
    def test_unknown_country_blocks_too(self, code):
        """Непроверенный exit не гарантирует Германию — это тоже отказ."""
        with pytest.raises(utils.ProxyCountryError) as excinfo:
            utils.require_german_exit(code, "proxy.host:8080")

        assert "страна не определена" in str(excinfo.value)

    def test_error_carries_parts_for_the_structured_log(self):
        with pytest.raises(utils.ProxyCountryError) as excinfo:
            utils.require_german_exit("CY", "user:s3cr3t@proxy.host:8080")

        error = excinfo.value
        assert error.reason.startswith("exit-IP не Германия: CY")
        assert error.proxy_address == "proxy.host:8080"


def test_get_domains_warns_when_blacklist_is_empty(set_paths, tmp_path, caplog):
    # Прод-инцидент: пустой domains.txt и без behavior.own_domain — чёрный
    # список выключен. Оператор должен увидеть это в логе, а не догадываться,
    # почему запрещённый домен кликается.
    domains_file = tmp_path / "empty_domains.txt"
    domains_file.write_text("\n  \n", "utf-8")
    set_paths(filtered_domains=domains_file)

    assert utils.get_domains() == []

    record = next(
        entry for entry in caplog.records if entry.getMessage().startswith("Blacklist is empty")
    )
    assert record.category == "click"
    assert record.levelname == "WARNING"
